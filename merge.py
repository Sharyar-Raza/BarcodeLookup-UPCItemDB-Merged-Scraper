import os
import random
import re
import sqlite3
import time

import pandas as pd
from bs4 import BeautifulSoup
from openpyxl import load_workbook

import undetected_chromedriver as uc
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException


# ============================================================
# CONFIG
# ============================================================

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

INPUT_EXCEL = os.path.join(
    SCRIPT_DIR,
    "Book1.xlsx"
)

# None = first/active sheet
# Agar specific sheet chahiye:
# SHEET_NAME = "Sheriyar"
SHEET_NAME = None


DB_NAME = os.path.join(
    SCRIPT_DIR,
    "merged_barcode_upc.db"
)

OUTPUT_EXCEL = os.path.join(
    SCRIPT_DIR,
    "merged_barcode_upc_output.xlsx"
)

CHROME_PROFILE_DIR = os.path.join(
    SCRIPT_DIR,
    "chrome_profile_merged_lookup"
)


PAGE_LOAD_TIMEOUT = 45
PAGE_RETRIES = 3
SKU_RETRIES = 3

VERIFICATION_TIMEOUT = 900

MIN_DELAY = 4.0
MAX_DELAY = 7.0

BETWEEN_SITES_MIN = 2.0
BETWEEN_SITES_MAX = 4.0

EXCEL_EXPORT_EVERY = 25

SHORT_BREAK_EVERY = 10
SHORT_BREAK_MIN = 15
SHORT_BREAK_MAX = 30

LONG_BREAK_EVERY = 25
LONG_BREAK_MIN = 60
LONG_BREAK_MAX = 120


BARCODELOOKUP_BASE = (
    "https://www.barcodelookup.com/"
)

UPCITEMDB_BASE = (
    "https://www.upcitemdb.com/"
)


UPC_ERROR_KEYWORDS = [

    "too fast",

    "429 too many requests",

    "access denied",

    "robot check",

    "api error",

    "limit error",

    "rate limit",

    "reach the limit",

    "take a break now",

]


UPC_NO_RESULT_PHRASES = [

    "no results found",

    "no result found",

    "no products found",

    "nothing found",

    "your search did not match",

    "couldn't find any",

    "could not find any",

    "0 results",

    "0 result",

]


# ============================================================
# HELPERS
# ============================================================

def human_delay(
    min_s=MIN_DELAY,
    max_s=MAX_DELAY
):

    delay = random.uniform(
        min_s,
        max_s
    )

    print(
        f"    [Delay] {delay:.1f}s",
        flush=True
    )

    time.sleep(
        delay
    )


def sku_key(
    value
):

    return str(
        value or ""
    ).strip().upper()


def normalize(
    value
):

    return re.sub(
        r"[^a-z0-9]",
        "",
        str(
            value or ""
        ).lower()
    )


def only_digits(
    value
):

    return re.sub(
        r"\D",
        "",
        str(
            value or ""
        )
    )


def get_body_text(
    driver
):

    try:

        return (
            driver
            .find_element(
                By.TAG_NAME,
                "body"
            )
            .text
            or ""
        )

    except Exception:

        return ""


def empty_result(
    part_number,
    status="Error",
    source="",
    url=""
):

    return {

        "part_number":
            str(
                part_number
            ).strip(),

        "status":
            status,

        "source":
            source,

        "product_title":
            "N/A",

        "upc":
            "N/A",

        "ean":
            "N/A",

        "brand":
            "N/A",

        "url":
            url,

        "barcode_status":
            "",

        "upcitemdb_status":
            "",

        "barcode_product_title":
            "",

        "barcode_upc":
            "",

        "barcode_ean":
            "",

        "barcode_brand":
            "",

        "barcode_url":
            "",

        "upcitemdb_product_title":
            "",

        "upcitemdb_upc":
            "",

        "upcitemdb_ean":
            "",

        "upcitemdb_brand":
            "",

        "upcitemdb_url":
            "",

    }


# ============================================================
# EXCEL
# ============================================================

def _excel_cell_to_sku(
    cell
):

    value = cell.value

    if value is None:

        return None

    if isinstance(
        value,
        str
    ):

        return value.strip()

    if (
        isinstance(
            value,
            (
                int,
                float
            )
        )
        and not isinstance(
            value,
            bool
        )
    ):

        number_format = str(
            cell.number_format
            or ""
        )

        if (
            isinstance(
                value,
                float
            )
            and value.is_integer()
        ):

            value = int(
                value
            )

        if isinstance(
            value,
            int
        ):

            zero_mask = re.fullmatch(
                r"0+",
                number_format
            )

            if zero_mask:

                return str(
                    value
                ).zfill(
                    len(
                        number_format
                    )
                )

        return str(
            value
        ).strip()

    return str(
        value
    ).strip()


def load_skus_from_excel(
    file_path=INPUT_EXCEL
):

    if not os.path.exists(
        file_path
    ):

        raise FileNotFoundError(
            f"Excel file nahi mili: {file_path}"
        )

    wb = load_workbook(
        file_path,
        data_only=True,
        read_only=False
    )

    try:

        if SHEET_NAME:

            if (
                SHEET_NAME
                not in wb.sheetnames
            ):

                raise ValueError(
                    f"Sheet '{SHEET_NAME}' nahi mili. "
                    f"Available: {wb.sheetnames}"
                )

            ws = wb[
                SHEET_NAME
            ]

        else:

            ws = wb.active

        sku_col = None

        for cell in ws[1]:

            header = str(
                cell.value
                or ""
            ).strip().lower()

            if header in {

                "sku",

                "part_number",

                "part number",

                "sku / part number",

            }:

                sku_col = cell.column
                break

        if sku_col is None:

            sku_col = 1

        skus = []

        for row in range(
            2,
            ws.max_row + 1
        ):

            sku = _excel_cell_to_sku(

                ws.cell(
                    row=row,
                    column=sku_col
                )

            )

            if sku:

                skus.append(
                    sku
                )

        print(

            f"[Excel] Loaded {len(skus)} SKUs from "
            f"'{os.path.basename(file_path)}' / "
            f"sheet '{ws.title}'.",

            flush=True

        )

        return skus

    finally:

        wb.close()


# ============================================================
# DATABASE
# ============================================================

def init_db():

    conn = sqlite3.connect(
        DB_NAME,
        timeout=30
    )

    try:

        cursor = conn.cursor()

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS products (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                part_number TEXT UNIQUE,

                status TEXT,

                source TEXT,

                product_title TEXT,

                upc TEXT,

                ean TEXT,

                brand TEXT,

                url TEXT,

                barcode_status TEXT,

                upcitemdb_status TEXT,

                barcode_product_title TEXT,

                barcode_upc TEXT,

                barcode_ean TEXT,

                barcode_brand TEXT,

                barcode_url TEXT,

                upcitemdb_product_title TEXT,

                upcitemdb_upc TEXT,

                upcitemdb_ean TEXT,

                upcitemdb_brand TEXT,

                upcitemdb_url TEXT,

                updated_at TEXT DEFAULT CURRENT_TIMESTAMP

            )
            """
        )

        # Existing DB ko bhi same schema par migrate karein.
        # Purana resume/data delete nahi hoga.

        cursor.execute(
            "PRAGMA table_info(products)"
        )

        existing_columns = {
            row[1]
            for row in cursor.fetchall()
        }

        extra_columns = {

            "barcode_product_title":
                "TEXT",

            "barcode_upc":
                "TEXT",

            "barcode_ean":
                "TEXT",

            "barcode_brand":
                "TEXT",

            "barcode_url":
                "TEXT",

            "upcitemdb_product_title":
                "TEXT",

            "upcitemdb_upc":
                "TEXT",

            "upcitemdb_ean":
                "TEXT",

            "upcitemdb_brand":
                "TEXT",

            "upcitemdb_url":
                "TEXT",

        }

        for (
            column_name,
            column_type
        ) in extra_columns.items():

            if (
                column_name
                not in existing_columns
            ):

                cursor.execute(
                    f"ALTER TABLE products "
                    f"ADD COLUMN "
                    f"{column_name} "
                    f"{column_type}"
                )

        # ----------------------------------------------------
        # BACKFILL OLD RECORDS
        # ----------------------------------------------------

        cursor.execute(
            """
            UPDATE products

            SET

                barcode_status = CASE

                    WHEN source = 'BarcodeLookup'
                    AND COALESCE(
                        barcode_status,
                        ''
                    ) = ''

                    THEN status

                    ELSE barcode_status

                END,


                barcode_product_title = CASE

                    WHEN source = 'BarcodeLookup'
                    AND COALESCE(
                        barcode_product_title,
                        ''
                    ) = ''

                    THEN product_title

                    ELSE barcode_product_title

                END,


                barcode_upc = CASE

                    WHEN source = 'BarcodeLookup'
                    AND COALESCE(
                        barcode_upc,
                        ''
                    ) = ''

                    THEN upc

                    ELSE barcode_upc

                END,


                barcode_ean = CASE

                    WHEN source = 'BarcodeLookup'
                    AND COALESCE(
                        barcode_ean,
                        ''
                    ) = ''

                    THEN ean

                    ELSE barcode_ean

                END,


                barcode_brand = CASE

                    WHEN source = 'BarcodeLookup'
                    AND COALESCE(
                        barcode_brand,
                        ''
                    ) = ''

                    THEN brand

                    ELSE barcode_brand

                END,


                barcode_url = CASE

                    WHEN source = 'BarcodeLookup'
                    AND COALESCE(
                        barcode_url,
                        ''
                    ) = ''

                    THEN url

                    ELSE barcode_url

                END,


                upcitemdb_status = CASE

                    WHEN source IN (
                        'UPCItemDB',
                        'Both'
                    )
                    AND COALESCE(
                        upcitemdb_status,
                        ''
                    ) = ''

                    THEN status

                    ELSE upcitemdb_status

                END,


                upcitemdb_product_title = CASE

                    WHEN source IN (
                        'UPCItemDB',
                        'Both'
                    )
                    AND COALESCE(
                        upcitemdb_product_title,
                        ''
                    ) = ''

                    THEN product_title

                    ELSE upcitemdb_product_title

                END,


                upcitemdb_upc = CASE

                    WHEN source IN (
                        'UPCItemDB',
                        'Both'
                    )
                    AND COALESCE(
                        upcitemdb_upc,
                        ''
                    ) = ''

                    THEN upc

                    ELSE upcitemdb_upc

                END,


                upcitemdb_ean = CASE

                    WHEN source IN (
                        'UPCItemDB',
                        'Both'
                    )
                    AND COALESCE(
                        upcitemdb_ean,
                        ''
                    ) = ''

                    THEN ean

                    ELSE upcitemdb_ean

                END,


                upcitemdb_brand = CASE

                    WHEN source IN (
                        'UPCItemDB',
                        'Both'
                    )
                    AND COALESCE(
                        upcitemdb_brand,
                        ''
                    ) = ''

                    THEN brand

                    ELSE upcitemdb_brand

                END,


                upcitemdb_url = CASE

                    WHEN source IN (
                        'UPCItemDB',
                        'Both'
                    )
                    AND COALESCE(
                        upcitemdb_url,
                        ''
                    ) = ''

                    THEN url

                    ELSE upcitemdb_url

                END
            """
        )

        conn.commit()

    finally:

        conn.close()


def get_completed_skus():

    conn = sqlite3.connect(
        DB_NAME,
        timeout=30
    )

    try:

        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT part_number
            FROM products
            WHERE status IN (
                'Found',
                'Not Found'
            )
            """
        )

        return {

            sku_key(
                row[0]
            )

            for row
            in cursor.fetchall()

            if row[0] is not None

        }

    finally:

        conn.close()


def save_to_db(
    result
):

    conn = sqlite3.connect(
        DB_NAME,
        timeout=30
    )

    try:

        cursor = conn.cursor()

        cursor.execute(
            """
            INSERT INTO products (

                part_number,

                status,

                source,

                product_title,

                upc,

                ean,

                brand,

                url,

                barcode_status,

                upcitemdb_status,

                barcode_product_title,

                barcode_upc,

                barcode_ean,

                barcode_brand,

                barcode_url,

                upcitemdb_product_title,

                upcitemdb_upc,

                upcitemdb_ean,

                upcitemdb_brand,

                upcitemdb_url,

                updated_at

            )

            VALUES (

                ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,

                ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,

                CURRENT_TIMESTAMP

            )

            ON CONFLICT(part_number)

            DO UPDATE SET

                status =
                    excluded.status,

                source =
                    excluded.source,

                product_title =
                    excluded.product_title,

                upc =
                    excluded.upc,

                ean =
                    excluded.ean,

                brand =
                    excluded.brand,

                url =
                    excluded.url,

                barcode_status =
                    excluded.barcode_status,

                upcitemdb_status =
                    excluded.upcitemdb_status,

                barcode_product_title =
                    excluded.barcode_product_title,

                barcode_upc =
                    excluded.barcode_upc,

                barcode_ean =
                    excluded.barcode_ean,

                barcode_brand =
                    excluded.barcode_brand,

                barcode_url =
                    excluded.barcode_url,

                upcitemdb_product_title =
                    excluded.upcitemdb_product_title,

                upcitemdb_upc =
                    excluded.upcitemdb_upc,

                upcitemdb_ean =
                    excluded.upcitemdb_ean,

                upcitemdb_brand =
                    excluded.upcitemdb_brand,

                upcitemdb_url =
                    excluded.upcitemdb_url,

                updated_at =
                    CURRENT_TIMESTAMP
            """,

            (

                result[
                    "part_number"
                ],

                result[
                    "status"
                ],

                result[
                    "source"
                ],

                result[
                    "product_title"
                ],

                result[
                    "upc"
                ],

                result[
                    "ean"
                ],

                result[
                    "brand"
                ],

                result[
                    "url"
                ],

                result[
                    "barcode_status"
                ],

                result[
                    "upcitemdb_status"
                ],

                result.get(
                    "barcode_product_title",
                    ""
                ),

                result.get(
                    "barcode_upc",
                    ""
                ),

                result.get(
                    "barcode_ean",
                    ""
                ),

                result.get(
                    "barcode_brand",
                    ""
                ),

                result.get(
                    "barcode_url",
                    ""
                ),

                result.get(
                    "upcitemdb_product_title",
                    ""
                ),

                result.get(
                    "upcitemdb_upc",
                    ""
                ),

                result.get(
                    "upcitemdb_ean",
                    ""
                ),

                result.get(
                    "upcitemdb_brand",
                    ""
                ),

                result.get(
                    "upcitemdb_url",
                    ""
                ),

            )

        )

        conn.commit()

    finally:

        conn.close()


def export_to_excel():

    try:

        conn = sqlite3.connect(
            DB_NAME
        )

        try:

            df = pd.read_sql_query(
                """
                SELECT

                    part_number
                        AS SKU,

                    status
                        AS final_status,

                    source
                        AS final_source,

                    product_title
                        AS final_product_title,

                    upc
                        AS final_upc,

                    ean
                        AS final_ean,

                    brand
                        AS final_brand,

                    url
                        AS final_url,


                    barcode_status,

                    barcode_product_title,

                    barcode_upc,

                    barcode_ean,

                    barcode_brand,

                    barcode_url,


                    upcitemdb_status,

                    upcitemdb_product_title,

                    upcitemdb_upc,

                    upcitemdb_ean,

                    upcitemdb_brand,

                    upcitemdb_url,


                    updated_at

                FROM products

                ORDER BY id
                """,

                conn

            )

        finally:

            conn.close()

        # ====================================================
        # SINGLE EXCEL:
        # FINAL + BARCODELOOKUP + UPCITEMDB
        # ====================================================

        df.to_excel(
            OUTPUT_EXCEL,
            index=False
        )

        print(

            f"[SAVE] "
            f"{os.path.basename(OUTPUT_EXCEL)} "
            f"updated ({len(df)} rows).",

            flush=True

        )

    except Exception as e:

        print(
            f"[Excel Save Error] {e}",
            flush=True
        )


# ============================================================
# CHROME
# ============================================================

def make_chrome_options():

    os.makedirs(
        CHROME_PROFILE_DIR,
        exist_ok=True
    )

    options = uc.ChromeOptions()

    options.add_argument(
        f"--user-data-dir="
        f"{CHROME_PROFILE_DIR}"
    )

    options.add_argument(
        "--window-size=1400,900"
    )

    options.add_argument(
        "--disable-notifications"
    )

    options.add_argument(
        "--disable-popup-blocking"
    )

    options.add_argument(
        "--disable-default-apps"
    )

    return options


def create_driver():

    print(
        "[Chrome] Starting persistent browser...",
        flush=True
    )

    try:

        driver = uc.Chrome(
            options=make_chrome_options()
        )

    except Exception as e:

        error_text = str(
            e
        )

        print(
            f"[Chrome] First start failed: "
            f"{error_text}",
            flush=True
        )

        match = re.search(
            r"Current browser version is "
            r"(\d+)\.",
            error_text
        )

        if not match:

            raise

        driver = uc.Chrome(

            options=
                make_chrome_options(),

            version_main=
                int(
                    match.group(1)
                )

        )

    driver.set_page_load_timeout(
        PAGE_LOAD_TIMEOUT
    )

    print(
        "[Chrome] Browser ready.",
        flush=True
    )

    return driver


# ============================================================
# SAFE GET
# ============================================================

def safe_get(
    driver,
    url
):

    for attempt in range(
        1,
        PAGE_RETRIES + 1
    ):

        try:

            driver.get(
                url
            )

            return True

        except TimeoutException:

            print(
                f"    [Timeout] "
                f"{attempt}/{PAGE_RETRIES}",
                flush=True
            )

            try:

                driver.execute_script(
                    "window.stop();"
                )

            except Exception:

                pass

            try:

                html = (
                    driver.page_source
                    or ""
                )

                if len(
                    html
                ) > 3000:

                    print(
                        "    Partial page loaded.",
                        flush=True
                    )

                    return True

            except Exception:

                pass

        except Exception as e:

            print(

                f"    [Network/Page Error] "
                f"{attempt}/{PAGE_RETRIES}: "
                f"{e}",

                flush=True

            )

        if attempt < PAGE_RETRIES:

            wait_seconds = (
                attempt * 10
            )

            print(
                f"    Retrying in "
                f"{wait_seconds}s...",
                flush=True
            )

            time.sleep(
                wait_seconds
            )

    return False


# ============================================================
# BARCODELOOKUP VERIFICATION
# ============================================================

def is_verification_page(
    driver
):

    try:

        title = (
            driver.title
            or ""
        ).strip().lower()

        body_text = (
            get_body_text(
                driver
            )
            .strip()
            .lower()
        )

        title_phrases = [

            "just a moment",

            "attention required",

            "security verification",

        ]

        body_phrases = [

            "verify you are human",

            "verifying you are human",

            "checking your browser",

            "performing security verification",

            "please verify you are a human",

        ]

        return (

            any(
                x in title
                for x
                in title_phrases
            )

            or

            any(
                x in body_text
                for x
                in body_phrases
            )

        )

    except Exception:

        return False


def wait_for_verification(
    driver
):

    if not is_verification_page(
        driver
    ):

        return True

    print(
        "[VERIFICATION DETECTED] Waiting...",
        flush=True
    )

    started = time.time()
    last_print = 0

    while True:

        elapsed = (
            time.time()
            - started
        )

        if (
            elapsed
            >= VERIFICATION_TIMEOUT
        ):

            print(
                "[Verification] Timeout.",
                flush=True
            )

            return False

        try:

            if not is_verification_page(
                driver
            ):

                print(
                    "[Verification] Completed. "
                    "Cooldown 10s...",
                    flush=True
                )

                time.sleep(
                    10
                )

                return True

        except Exception:

            pass

        if (
            time.time()
            - last_print
            >= 15
        ):

            print(
                "[Verification] Waiting...",
                flush=True
            )

            last_print = time.time()

        time.sleep(
            2
        )


# ============================================================
# BARCODELOOKUP EXTRACTION
# ============================================================

def extract_barcodelookup_data(
    driver
):

    title = "N/A"
    upc = "N/A"
    ean = "N/A"
    brand = "N/A"

    try:

        title_elem = (
            driver
            .find_element(
                By.TAG_NAME,
                "h1"
            )
        )

        text = (
            title_elem.text
            .strip()
        )

        if text:

            title = text

    except Exception:

        pass

    soup = BeautifulSoup(
        driver.page_source,
        "html.parser"
    )

    full_text = soup.get_text(
        " ",
        strip=True
    )

    ean_match = re.search(

        r"\bEAN(?:-13)?"
        r"\s*[:\-]?\s*"
        r"(\d{12,13})\b",

        full_text,

        re.IGNORECASE

    )

    if ean_match:

        ean = ean_match.group(
            1
        )

    upc_match = re.search(

        r"\bUPC(?:-A)?"
        r"\s*[:\-]?\s*"
        r"(\d{12,13})\b",

        full_text,

        re.IGNORECASE

    )

    if upc_match:

        upc = upc_match.group(
            1
        )

    brand_match = re.search(

        r"Brand\s*[:\-]?\s*"
        r"([A-Za-z0-9\s.\-&]+?)"
        r"(?=\s*(?:"
        r"Category|Manufacturer|Model|UPC|EAN|$"
        r"))",

        full_text,

        re.IGNORECASE

    )

    if brand_match:

        brand = (
            brand_match
            .group(1)
            .strip()
        )

    return (
        title,
        upc,
        ean,
        brand,
        full_text
    )


def barcodelookup_says_not_found(
    title,
    full_text
):

    combined = (
        f"{title} "
        f"{full_text}"
    ).lower()

    phrases = [

        "doesn't exist",

        "does not exist",

        "product not found",

        "no product found",

        "no products found",

        "we couldn't find",

        "we could not find",

    ]

    return any(
        phrase in combined
        for phrase
        in phrases
    )


def find_barcodelookup_product_url(
    driver
):

    try:

        soup = BeautifulSoup(
            driver.page_source,
            "html.parser"
        )

        for link in soup.select(
            "a[href]"
        ):

            href = (
                link.get(
                    "href",
                    ""
                )
                or ""
            ).strip()

            if re.search(
                r"/\d{8,14}/?$",
                href
            ):

                if href.startswith(
                    "http"
                ):

                    return href

                if href.startswith(
                    "/"
                ):

                    return (
                        "https://www.barcodelookup.com"
                        + href
                    )

    except Exception:

        pass

    return None


def scrape_barcodelookup_once(
    driver,
    part_number
):

    part_number = str(
        part_number
    ).strip()

    search_url = (
        f"{BARCODELOOKUP_BASE}"
        f"{part_number}"
    )

    result = empty_result(

        part_number,

        status="Error",

        source="BarcodeLookup",

        url=search_url

    )

    result[
        "barcode_status"
    ] = "Error"

    print(
        f"    [1/2 BarcodeLookup] "
        f"{search_url}",
        flush=True
    )

    if not safe_get(
        driver,
        search_url
    ):

        return result

    if not wait_for_verification(
        driver
    ):

        return result

    human_delay()

    try:

        (
            title,
            upc,
            ean,
            brand,
            full_text
        ) = extract_barcodelookup_data(
            driver
        )

    except Exception as e:

        print(
            f"    [Barcode extraction error] "
            f"{e}",
            flush=True
        )

        return result

    if (
        upc != "N/A"
        or ean != "N/A"
    ):

        result.update(
            {

                "status":
                    "Found",

                "source":
                    "BarcodeLookup",

                "product_title":
                    title,

                "upc":
                    upc,

                "ean":
                    ean,

                "brand":
                    brand,

                "url":
                    driver.current_url,

                "barcode_status":
                    "Found",

            }
        )

        return result

    if barcodelookup_says_not_found(
        title,
        full_text
    ):

        result[
            "status"
        ] = "Not Found"

        result[
            "barcode_status"
        ] = "Not Found"

        result[
            "product_title"
        ] = title

        result[
            "url"
        ] = driver.current_url

        return result

    target_product_url = (
        find_barcodelookup_product_url(
            driver
        )
    )

    if (
        target_product_url
        and
        target_product_url
        != driver.current_url
    ):

        print(
            "    Opening BarcodeLookup "
            "product detail page...",
            flush=True
        )

        if not safe_get(
            driver,
            target_product_url
        ):

            return result

        if not wait_for_verification(
            driver
        ):

            return result

        human_delay()

        try:

            (
                title,
                upc,
                ean,
                brand,
                full_text
            ) = extract_barcodelookup_data(
                driver
            )

        except Exception as e:

            print(
                f"    [Barcode detail "
                f"extraction error] {e}",
                flush=True
            )

            return result

        if (
            upc != "N/A"
            or ean != "N/A"
        ):

            result.update(
                {

                    "status":
                        "Found",

                    "source":
                        "BarcodeLookup",

                    "product_title":
                        title,

                    "upc":
                        upc,

                    "ean":
                        ean,

                    "brand":
                        brand,

                    "url":
                        driver.current_url,

                    "barcode_status":
                        "Found",

                }
            )

            return result

        if barcodelookup_says_not_found(
            title,
            full_text
        ):

            result[
                "status"
            ] = "Not Found"

            result[
                "barcode_status"
            ] = "Not Found"

            result[
                "product_title"
            ] = title

            result[
                "url"
            ] = driver.current_url

            return result

    result[
        "status"
    ] = "Unclear"

    result[
        "barcode_status"
    ] = "Unclear"

    return result


def process_barcodelookup_with_retries(
    driver,
    sku
):

    result = None

    for attempt in range(
        1,
        SKU_RETRIES + 1
    ):

        print(
            f"    Barcode attempt "
            f"{attempt}/{SKU_RETRIES}",
            flush=True
        )

        result = scrape_barcodelookup_once(
            driver,
            sku
        )

        if result[
            "status"
        ] in {
            "Found",
            "Not Found"
        }:

            return result

        print(

            f"    [Barcode WARNING] "
            f"Status: "
            f"{result['status']}",

            flush=True

        )

        if attempt < SKU_RETRIES:

            wait_seconds = (
                attempt * 20
            )

            print(
                f"    Retrying BarcodeLookup "
                f"in {wait_seconds}s...",
                flush=True
            )

            time.sleep(
                wait_seconds
            )

    return result


# ============================================================
# UPCITEMDB
# ============================================================

def detect_upcitemdb_error(
    driver
):

    try:

        src = (
            driver.page_source
            or ""
        ).lower()

    except Exception:

        src = ""

    visible = get_body_text(
        driver
    ).lower()

    for keyword in UPC_ERROR_KEYWORDS:

        if (
            keyword in src
            or keyword in visible
        ):

            return keyword

    return None


def search_upcitemdb_exact_sku(
    driver,
    part_number
):

    exact_sku = str(
        part_number
    ).strip()

    print(
        f"    [2/2 UPCItemDB SEARCH] "
        f"{repr(exact_sku)}",
        flush=True
    )

    driver.get(
        UPCITEMDB_BASE
    )

    time.sleep(
        1
    )

    error = detect_upcitemdb_error(
        driver
    )

    if error:

        raise RuntimeError(
            f"SEARCH_BLOCK::{error}"
        )

    search_box = None

    selectors = [

        "input[name='q']",

        "input[type='search']",

        "form input[type='text']",

        "input.form-control",

    ]

    for selector in selectors:

        try:

            elements = driver.find_elements(
                By.CSS_SELECTOR,
                selector
            )

            for element in elements:

                try:

                    if (
                        element.is_displayed()
                        and
                        element.is_enabled()
                    ):

                        search_box = element
                        break

                except Exception:

                    continue

            if search_box:

                break

        except Exception:

            continue

    if search_box is None:

        raise RuntimeError(
            "SEARCH_BAR_NOT_FOUND"
        )

    search_box.click()

    search_box.send_keys(
        Keys.CONTROL,
        "a"
    )

    search_box.send_keys(
        exact_sku
    )

    print(
        f"    [TYPED] "
        f"{exact_sku}",
        flush=True
    )

    search_box.send_keys(
        Keys.ENTER
    )

    time.sleep(
        2
    )

    error = detect_upcitemdb_error(
        driver
    )

    if error:

        raise RuntimeError(
            f"SEARCH_BLOCK::{error}"
        )

    return driver.current_url


def find_verified_upcitemdb_result(
    driver,
    part_number
):

    exact_sku = str(
        part_number
    ).strip()

    exact_lower = (
        exact_sku.lower()
    )

    query_norm = normalize(
        exact_sku
    )

    try:

        WebDriverWait(
            driver,
            10
        ).until(

            lambda d: (

                len(
                    d.find_elements(
                        By.CSS_SELECTOR,
                        "a[href*='/upc/']"
                    )
                ) > 0

                or

                any(

                    phrase
                    in get_body_text(
                        d
                    ).lower()

                    for phrase
                    in UPC_NO_RESULT_PHRASES

                )

            )

        )

    except Exception:

        pass

    body_text = get_body_text(
        driver
    )

    body_lower = (
        body_text.lower()
    )

    if any(

        phrase in body_lower

        for phrase
        in UPC_NO_RESULT_PHRASES

    ):

        print(
            "    [UPCItemDB NOT FOUND] "
            "no-result message.",
            flush=True
        )

        return None

    try:

        raw_candidates = (
            driver.execute_script(
                """
                const anchors = Array.from(
                    document.querySelectorAll(
                        "a[href*='/upc/']"
                    )
                );

                const output = [];

                const seen = new Set();

                for (const a of anchors) {

                    const href =
                        a.href || "";

                    if (
                        !/\\/upc\\/\\d+(?:[/?#]|$)/i
                        .test(href)
                    ) {
                        continue;
                    }

                    let node = a;

                    for (
                        let level = 0;
                        level < 7 && node;
                        level++
                    ) {

                        const text =
                            (node.innerText || "")
                            .trim();

                        if (
                            text
                            && text.length <= 2000
                        ) {

                            const key =
                                href + "||" + text;

                            if (!seen.has(key)) {

                                output.push({
                                    href,
                                    text,
                                    level
                                });

                                seen.add(key);
                            }
                        }

                        node =
                            node.parentElement;
                    }
                }

                return output;
                """
            )
            or []
        )

    except Exception:

        raw_candidates = []

    matches = []

    for item in raw_candidates:

        try:

            href = str(
                item.get(
                    "href",
                    ""
                )
            )

            context = " ".join(
                str(
                    item.get(
                        "text",
                        ""
                    )
                ).split()
            )

            level = int(
                item.get(
                    "level",
                    99
                )
            )

            m = re.search(
                r"/upc/(\d+)"
                r"(?:[/?#]|$)",
                href,
                flags=re.I
            )

            if not m:

                continue

            barcode = m.group(
                1
            )

            context_lower = (
                context.lower()
            )

            context_norm = normalize(
                context
            )

            exact_match = (

                bool(
                    exact_lower
                )

                and

                exact_lower
                in context_lower

            )

            normalized_match = (

                len(
                    query_norm
                ) >= 6

                and

                query_norm
                in context_norm

            )

            if (
                not exact_match
                and
                not normalized_match
            ):

                continue

            score = 0

            if exact_match:

                score += 1000

            if normalized_match:

                score += 200

            score += max(
                0,
                100 - (
                    level * 10
                )
            )

            score += max(
                0,
                100
                - min(
                    len(
                        context
                    ),
                    1000
                )
                // 10
            )

            matches.append(
                {

                    "url":
                        (
                            "https://www.upcitemdb.com/"
                            f"upc/{barcode}"
                        ),

                    "barcode":
                        barcode,

                    "context":
                        context,

                    "score":
                        score,

                    "level":
                        level,

                }
            )

        except Exception:

            continue

    if not matches:

        print(

            f"    [UPCItemDB NOT FOUND] "
            f"'{exact_sku}' verified "
            f"result nahi mila.",

            flush=True

        )

        return None

    best_by_barcode = {}

    for item in matches:

        barcode = item[
            "barcode"
        ]

        if (

            barcode
            not in best_by_barcode

            or

            item[
                "score"
            ]
            >
            best_by_barcode[
                barcode
            ][
                "score"
            ]

        ):

            best_by_barcode[
                barcode
            ] = item

    unique_matches = list(
        best_by_barcode.values()
    )

    unique_matches.sort(

        key=lambda x: (

            -x[
                "score"
            ],

            x[
                "level"
            ],

            len(
                x[
                    "context"
                ]
            )

        )

    )

    best = unique_matches[
        0
    ]

    print(

        f"    [UPCItemDB MATCH] "
        f"SKU '{exact_sku}' "
        f"=> UPC "
        f"{best['barcode']}",

        flush=True

    )

    return best


def equivalent_barcode(
    value,
    expected
):

    a = only_digits(
        value
    )

    b = only_digits(
        expected
    )

    if (
        not a
        or not b
    ):

        return False

    return (

        a == b

        or

        a.lstrip(
            "0"
        )
        ==
        b.lstrip(
            "0"
        )

    )


def extract_upcitemdb_detail(
    driver,
    part_number,
    detail_url,
    expected_barcode
):

    result = empty_result(

        part_number,

        status="Not Found",

        source="UPCItemDB",

        url=detail_url

    )

    result[
        "upcitemdb_status"
    ] = "Not Found"

    wait = WebDriverWait(
        driver,
        5
    )

    try:

        title_elem = wait.until(

            EC.presence_of_element_located(
                (
                    By.CSS_SELECTOR,
                    "p.detailtitle, h1"
                )
            )

        )

        title = " ".join(
            title_elem.text.split()
        )

        if title:

            result[
                "product_title"
            ] = title

    except Exception:

        pass

    try:

        rows = driver.find_elements(

            By.XPATH,

            '//*[@id="info"]'
            '//table//tr'

        )

        for row in rows:

            txt = " ".join(
                (
                    row.text
                    or ""
                ).split()
            )

            if not txt:

                continue

            upc_match = re.search(

                r"\bUPC(?:-A)?"
                r"\s*:\s*"
                r"([0-9][0-9 \-]{5,20}[0-9])",

                txt,

                flags=re.I

            )

            if upc_match:

                value = only_digits(
                    upc_match.group(
                        1
                    )
                )

                if value:

                    result[
                        "upc"
                    ] = value

            ean_match = re.search(

                r"\bEAN(?:-13)?"
                r"\s*:\s*"
                r"([0-9][0-9 \-]{5,20}[0-9])",

                txt,

                flags=re.I

            )

            if ean_match:

                value = only_digits(
                    ean_match.group(
                        1
                    )
                )

                if value:

                    result[
                        "ean"
                    ] = value

            brand_match = re.search(

                r"\bBrand\s*:\s*(.+)$",

                txt,

                flags=re.I

            )

            if brand_match:

                brand = (
                    brand_match
                    .group(1)
                    .strip()
                )

                if brand:

                    result[
                        "brand"
                    ] = brand

    except Exception:

        pass

    barcode_verified = (

        equivalent_barcode(
            result[
                "upc"
            ],
            expected_barcode
        )

        or

        equivalent_barcode(
            result[
                "ean"
            ],
            expected_barcode
        )

    )

    if not barcode_verified:

        print(

            f"    [UPCItemDB VERIFY FAILED] "
            f"{part_number}: "
            f"URL barcode "
            f"{expected_barcode} "
            f"!= extracted UPC/EAN",

            flush=True

        )

        return result

    result[
        "status"
    ] = "Found"

    result[
        "upcitemdb_status"
    ] = "Found"

    print(

        f"    [UPCItemDB VERIFIED FOUND] "
        f"SKU: {part_number} | "
        f"UPC: {result['upc']} | "
        f"EAN: {result['ean']}",

        flush=True

    )

    return result


def scrape_upcitemdb_once(
    driver,
    part_number
):

    exact_sku = str(
        part_number
    ).strip()

    search_url = (

        "https://www.upcitemdb.com/"
        f"s?q={exact_sku}"

    )

    result = empty_result(

        exact_sku,

        status="Error",

        source="UPCItemDB",

        url=search_url

    )

    result[
        "upcitemdb_status"
    ] = "Error"

    try:

        search_url = (
            search_upcitemdb_exact_sku(
                driver,
                exact_sku
            )
        )

        error = detect_upcitemdb_error(
            driver
        )

        if error:

            raise RuntimeError(
                f"SEARCH_BLOCK::{error}"
            )

        verified = (
            find_verified_upcitemdb_result(
                driver,
                exact_sku
            )
        )

        if not verified:

            result[
                "status"
            ] = "Not Found"

            result[
                "upcitemdb_status"
            ] = "Not Found"

            result[
                "url"
            ] = search_url

            return result

        detail_url = verified[
            "url"
        ]

        expected_barcode = verified[
            "barcode"
        ]

        print(
            f"    [OPEN UPCItemDB PRODUCT] "
            f"{detail_url}",
            flush=True
        )

        driver.get(
            detail_url
        )

        time.sleep(
            1
        )

        error = detect_upcitemdb_error(
            driver
        )

        if error:

            raise RuntimeError(
                f"DETAIL_BLOCK::{error}"
            )

        return extract_upcitemdb_detail(

            driver,

            exact_sku,

            detail_url,

            expected_barcode

        )

    except Exception as e:

        print(
            f"    [UPCItemDB TECHNICAL ERROR] "
            f"{e}",
            flush=True
        )

        result[
            "status"
        ] = "Error"

        result[
            "upcitemdb_status"
        ] = "Error"

        return result


def process_upcitemdb_with_retries(
    driver,
    sku
):

    result = None

    for attempt in range(
        1,
        SKU_RETRIES + 1
    ):

        print(

            f"    UPCItemDB attempt "
            f"{attempt}/{SKU_RETRIES}",

            flush=True

        )

        result = scrape_upcitemdb_once(
            driver,
            sku
        )

        if result[
            "status"
        ] in {
            "Found",
            "Not Found"
        }:

            return result

        print(

            f"    [UPCItemDB WARNING] "
            f"Status: "
            f"{result['status']}",

            flush=True

        )

        if attempt < SKU_RETRIES:

            wait_seconds = (
                attempt * 20
            )

            print(

                f"    Retrying UPCItemDB "
                f"in {wait_seconds}s...",

                flush=True

            )

            time.sleep(
                wait_seconds
            )

    return result


# ============================================================
# KEEP BOTH WEBSITE RESULTS IN SAME FINAL RECORD
# ============================================================

def attach_barcode_output(
    target,
    barcode_result
):

    target[
        "barcode_status"
    ] = barcode_result.get(
        "status",
        ""
    )

    target[
        "barcode_product_title"
    ] = barcode_result.get(
        "product_title",
        ""
    )

    target[
        "barcode_upc"
    ] = barcode_result.get(
        "upc",
        ""
    )

    target[
        "barcode_ean"
    ] = barcode_result.get(
        "ean",
        ""
    )

    target[
        "barcode_brand"
    ] = barcode_result.get(
        "brand",
        ""
    )

    target[
        "barcode_url"
    ] = barcode_result.get(
        "url",
        ""
    )

    return target


def attach_upcitemdb_output(
    target,
    upc_result
):

    target[
        "upcitemdb_status"
    ] = upc_result.get(
        "status",
        ""
    )

    target[
        "upcitemdb_product_title"
    ] = upc_result.get(
        "product_title",
        ""
    )

    target[
        "upcitemdb_upc"
    ] = upc_result.get(
        "upc",
        ""
    )

    target[
        "upcitemdb_ean"
    ] = upc_result.get(
        "ean",
        ""
    )

    target[
        "upcitemdb_brand"
    ] = upc_result.get(
        "brand",
        ""
    )

    target[
        "upcitemdb_url"
    ] = upc_result.get(
        "url",
        ""
    )

    return target


# ============================================================
# MERGED SKU FLOW
# ============================================================

def process_one_sku(
    driver,
    sku
):

    sku = str(
        sku
    ).strip()

    print()

    print(
        "=" * 64
    )

    print(
        f"SKU: {sku}"
    )

    print(
        "FLOW: BarcodeLookup "
        "-> if Not Found "
        "-> UPCItemDB"
    )

    print(
        "=" * 64
    )

    # ========================================================
    # FIRST: BARCODELOOKUP
    # ========================================================

    barcode_result = (
        process_barcodelookup_with_retries(
            driver,
            sku
        )
    )

    # --------------------------------------------------------
    # FOUND ON BARCODELOOKUP
    # --------------------------------------------------------

    if (
        barcode_result[
            "status"
        ]
        == "Found"
    ):

        barcode_result[
            "barcode_status"
        ] = "Found"

        attach_barcode_output(
            barcode_result,
            barcode_result
        )

        barcode_result[
            "upcitemdb_status"
        ] = "Skipped"

        barcode_result[
            "upcitemdb_product_title"
        ] = ""

        barcode_result[
            "upcitemdb_upc"
        ] = ""

        barcode_result[
            "upcitemdb_ean"
        ] = ""

        barcode_result[
            "upcitemdb_brand"
        ] = ""

        barcode_result[
            "upcitemdb_url"
        ] = ""

        print(

            f"    [SKIP UPCItemDB] "
            f"BarcodeLookup par "
            f"SKU '{sku}' mil gaya.",

            flush=True

        )

        return barcode_result

    # --------------------------------------------------------
    # ERROR / UNCLEAR ON BARCODELOOKUP
    # --------------------------------------------------------

    if (
        barcode_result[
            "status"
        ]
        != "Not Found"
    ):

        attach_barcode_output(
            barcode_result,
            barcode_result
        )

        barcode_result[
            "barcode_status"
        ] = barcode_result[
            "status"
        ]

        barcode_result[
            "upcitemdb_status"
        ] = "Not Run"

        barcode_result[
            "status"
        ] = "Error"

        barcode_result[
            "source"
        ] = "BarcodeLookup"

        print(

            "    [STOP CURRENT SKU] "
            "BarcodeLookup ka result "
            "reliable Not Found nahi tha. "
            "UPCItemDB fallback nahi chalega. "
            "Next run mein retry hoga.",

            flush=True

        )

        return barcode_result

    # ========================================================
    # BARCODELOOKUP = NOT FOUND
    # NOW UPCITEMDB
    # ========================================================

    print(

        f"    [FALLBACK] "
        f"BarcodeLookup par "
        f"'{sku}' Not Found. "
        f"UPCItemDB search start...",

        flush=True

    )

    time.sleep(

        random.uniform(
            BETWEEN_SITES_MIN,
            BETWEEN_SITES_MAX
        )

    )

    upc_result = (
        process_upcitemdb_with_retries(
            driver,
            sku
        )
    )

    # Save BarcodeLookup raw result too.
    attach_barcode_output(
        upc_result,
        barcode_result
    )

    # Save UPCItemDB raw result too.
    upc_raw_copy = (
        upc_result.copy()
    )

    attach_upcitemdb_output(
        upc_result,
        upc_raw_copy
    )

    upc_result[
        "barcode_status"
    ] = "Not Found"

    # --------------------------------------------------------
    # UPCITEMDB FOUND
    # --------------------------------------------------------

    if (
        upc_result[
            "status"
        ]
        == "Found"
    ):

        upc_result[
            "source"
        ] = "UPCItemDB"

        return upc_result

    # --------------------------------------------------------
    # BOTH NOT FOUND
    # --------------------------------------------------------

    if (
        upc_result[
            "status"
        ]
        == "Not Found"
    ):

        upc_result[
            "source"
        ] = "Both"

        return upc_result

    # --------------------------------------------------------
    # UPCITEMDB ERROR
    # --------------------------------------------------------

    upc_result[
        "status"
    ] = "Error"

    upc_result[
        "source"
    ] = "UPCItemDB"

    return upc_result


# ============================================================
# MAIN
# ============================================================

def main():

    print()

    print(
        "============================================================"
    )

    print(
        "=== MERGED BARCODELOOKUP + UPCITEMDB WORKFLOW STARTED ==="
    )

    print(
        "============================================================"
    )

    init_db()

    try:

        all_skus = load_skus_from_excel(
            INPUT_EXCEL
        )

    except Exception as e:

        print(
            f"[STOP] Excel load error: {e}",
            flush=True
        )

        return

    if not all_skus:

        print(
            "[STOP] No SKUs found.",
            flush=True
        )

        return

    completed = get_completed_skus()

    remaining_skus = [

        sku

        for sku in all_skus

        if sku_key(
            sku
        ) not in completed

    ]

    already_completed = (

        len(
            all_skus
        )

        -

        len(
            remaining_skus
        )

    )

    print(
        f"Total SKUs        : "
        f"{len(all_skus)}"
    )

    print(
        f"Already completed : "
        f"{already_completed}"
    )

    print(
        f"Remaining         : "
        f"{len(remaining_skus)}"
    )

    print(
        f"Database          : "
        f"{DB_NAME}"
    )

    print(
        f"Output Excel      : "
        f"{OUTPUT_EXCEL}"
    )

    if not remaining_skus:

        print(
            "[DONE] Sab SKUs already completed hain."
        )

        export_to_excel()

        return

    driver = None
    new_saved = 0

    try:

        driver = create_driver()

        for (
            current_no,
            sku
        ) in enumerate(
            remaining_skus,
            start=1
        ):

            print()

            print(

                f"[Remaining "
                f"{current_no}/"
                f"{len(remaining_skus)} | "

                f"Overall "
                f"{already_completed + current_no}/"
                f"{len(all_skus)}]"

            )

            result = process_one_sku(
                driver,
                sku
            )

            # =================================================
            # ERROR => DO NOT SAVE
            # NEXT RUN RETRY SAME SKU
            # =================================================

            if result[
                "status"
            ] not in {
                "Found",
                "Not Found"
            }:

                print()

                print(
                    "!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!"
                )

                print(
                    f"[STOP] SKU {sku} "
                    f"complete nahi hui."
                )

                print(
                    f"Status: "
                    f"{result['status']}"
                )

                print(
                    "Ye SKU DB mein save nahi ki gayi."
                )

                print(
                    "Next run mein isi SKU ko "
                    "dobara attempt kiya jayega."
                )

                print(
                    "!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!"
                )

                break

            # =================================================
            # FINAL SAVE
            # =================================================

            save_to_db(
                result
            )

            new_saved += 1

            print(

                f"    FINAL STATUS : "
                f"{result['status']}\n"

                f"    SOURCE       : "
                f"{result['source']}\n"

                f"    BARCODE      : "
                f"{result['barcode_status']}\n"

                f"    UPCITEMDB    : "
                f"{result['upcitemdb_status']}\n"

                f"    UPC          : "
                f"{result['upc']}\n"

                f"    EAN          : "
                f"{result['ean']}\n"

                f"    Brand        : "
                f"{result['brand']}\n"

                f"    [DB] Saved successfully."

            )

            # =================================================
            # PERIODIC EXCEL
            # =================================================

            if (
                new_saved
                % EXCEL_EXPORT_EVERY
                == 0
            ):

                export_to_excel()

            # =================================================
            # BREAKS
            # =================================================

            if (
                new_saved
                % LONG_BREAK_EVERY
                == 0
            ):

                break_time = random.randint(
                    LONG_BREAK_MIN,
                    LONG_BREAK_MAX
                )

                print(
                    f"[LONG BREAK] "
                    f"Waiting {break_time}s...",
                    flush=True
                )

                time.sleep(
                    break_time
                )

            elif (
                new_saved
                % SHORT_BREAK_EVERY
                == 0
            ):

                break_time = random.randint(
                    SHORT_BREAK_MIN,
                    SHORT_BREAK_MAX
                )

                print(
                    f"[SHORT BREAK] "
                    f"Waiting {break_time}s...",
                    flush=True
                )

                time.sleep(
                    break_time
                )

            else:

                human_delay()

    except KeyboardInterrupt:

        print(
            "\n[STOPPED] Ctrl+C pressed. "
            "Saved DB records safe hain."
        )

    except Exception as e:

        print(
            f"\n[FATAL ERROR] "
            f"{repr(e)}"
        )

        print(
            "Database mein already saved "
            "records safe hain."
        )

    finally:

        print(
            "\n[SAVE] Final Excel export...",
            flush=True
        )

        export_to_excel()

        if driver is not None:

            try:

                driver.quit()

            except Exception:

                pass

    print()

    print(
        "============================================================"
    )

    print(
        "[SESSION FINISHED]"
    )

    print(
        f"New records this run : "
        f"{new_saved}"
    )

    print(
        f"Database             : "
        f"{DB_NAME}"
    )

    print(
        f"Excel                : "
        f"{OUTPUT_EXCEL}"
    )

    print(
        "============================================================"
    )


if __name__ == "__main__":

    main()